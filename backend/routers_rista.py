"""Rista (DotPe) data loaded from files downloaded in the Rista portal, until the DotPe API is available.

Three Rista exports are understood:
- sales:     Reports > Sales > Sales Items > "Sales Items By Transaction" (one line per dish, option or add-on)
- cancelled: Reports > Sales > Sale Transactions > "Cancelled Transactions"
- menu:      Business Settings > Items > three dots (⋮) > "Export Items" (emailed by Rista)

Only the columns listed below are kept. Customer names, phone numbers, ids, tax ids, notes and staff names
in the sales lines are dropped on import (decision "Customer details" in docs/open-questions.md).
Sales and cancelled files replace what was stored for the same branch and business day, so loading the
same or an overlapping file twice never double counts.
"""
import asyncio
import csv
import io
import math
import os
import re
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell import WriteOnlyCell
from pymongo import ReplaceOne

from db import db, uid, now_iso, clean_list, audit
from auth import require_roles

router = APIRouter(prefix='/rista', tags=['rista'])

VIEW_ROLES = ('admin', 'management')
MAX_BYTES = 40 * 1024 * 1024
# A download holds every matching line in memory while the file is built; the hosting plan has little memory.
EXPORT_MAX_LINES = int(os.environ.get('RISTA_EXPORT_MAX_LINES', '100000'))

# Reports only change when an admin loads a file, so answers are kept in memory until then (the app runs as one
# process). Anything that can change outside a file load, such as recipes, is always read fresh.
_cache = {}


def _invalidate():
    _cache.clear()


async def _cached(key, compute):
    if key in _cache:
        return _cache[key]
    value = await compute()
    if len(_cache) >= 300:
        _cache.clear()
    _cache[key] = value
    return value

# (field we store, Rista column, kind) kind: s = text, n = number
SALES_COLUMNS = [
    ('branch_name', 'Branch Name', 's'), ('branch_code', 'Branch Code', 's'),
    ('business_brand', 'Business Brand', 's'), ('business_date', 'Business Date', 's'),
    ('invoice_date', 'Invoice Date', 's'), ('invoice_number', 'Invoice Number', 's'),
    ('order_source', 'Order Source', 's'), ('source_order_number', 'Source Order Number', 's'),
    ('invoice_type', 'Invoice Type', 's'), ('sale_status', 'Sale Status', 's'),
    ('sku', 'SKU', 's'), ('line_type', 'Type', 's'), ('item_name', 'Item Name', 's'), ('tags', 'Tags', 's'),
    ('parent_sku', 'Parent SKU (Sold With)', 's'), ('qty', 'Sale Item Qty', 'n'), ('unit_rate', 'Unit Rate', 'n'),
    ('gross', 'Gross Amount', 'n'), ('discounts', 'Discounts', 'n'), ('net', 'Net Amount', 'n'),
    ('taxes', 'Taxes', 'n'), ('aggregator_tax', 'Tax Collected by Aggregator', 'n'),
    ('direct_charge', 'Direct Charge', 'n'), ('indirect_charge', 'Indirect Charge', 'n'),
    ('materials_cost', 'Materials Cost', 'n'), ('channel', 'Channel', 's'), ('category', 'Category', 's'),
    ('subcategory', 'Subcategory', 's'), ('account_name', 'Account Name', 's'), ('variants', 'Variant(s)', 's'),
    ('time_to_prepare', 'Time To Prepare', 'n'), ('time_to_dispatch', 'Time To Dispatch', 'n'),
    ('item_discount_codes', 'Item Discount Codes', 's'), ('bill_discount_codes', 'Bill Discount Codes', 's'),
    ('payment_modes', 'Payment Modes', 's'),
]
CANCELLED_COLUMNS = [
    ('branch_name', 'Branch Name', 's'), ('branch_code', 'Branch Code', 's'),
    ('business_brand', 'Business Brand', 's'), ('business_date', 'Business Date', 's'),
    ('invoice_number', 'Invoice Number', 's'), ('order_number', 'Order Number', 's'),
    ('order_source', 'Order Source', 's'), ('source_order_number', 'Source Order Number', 's'),
    ('order_date', 'Order/Invoice Date', 's'), ('status', 'Status', 's'),
    ('fulfillment_status', 'Fulfillment Status', 's'), ('channel', 'Channel', 's'),
    ('time_to_cancel', 'Time to Cancel', 'n'), ('cancel_date', 'Cancel/Void Date', 's'),
    ('cancelled_by', 'Cancel/Void By', 's'), ('reason', 'Reason For Cancel/Void', 's'),
    ('source_reason', 'Source Reason for Cancel/Void', 's'),
    ('gross', 'Gross Amount', 'n'), ('discounts', 'Discounts', 'n'), ('net', 'Net Amount', 'n'),
    ('taxes', 'Taxes', 'n'), ('total', 'Total', 'n'), ('payment_modes', 'Payment Modes', 's'),
]
MENU_COLUMNS = [
    ('sku', 'SKU', 's'), ('category', 'CATEGORY', 's'), ('name', 'ITEM NAME', 's'),
    ('variant', 'VARIANT ATTRIBUTES', 's'), ('item_brand', 'ITEM BRAND', 's'), ('status', 'STATUS', 's'),
    ('subcategory', 'SUB CATEGORY', 's'), ('account', 'ACCOUNT', 's'), ('unit', 'MEASURING UNIT', 's'),
    ('price_includes_tax', 'PRICE INCLUDES TAX', 's'), ('mrp', 'MRP', 'n'), ('tags', 'TAGS', 's'),
    ('short_name', 'SHORT NAME', 's'), ('description', 'DESCRIPTION', 's'), ('options', 'OPTIONSET NAME', 's'),
]
# Columns a file must have to be accepted as that report (catches a wrong file early).
REQUIRED = {
    'sales': ['Branch Code', 'Business Date', 'Invoice Number', 'SKU', 'Type', 'Net Amount', 'Channel', 'Account Name'],
    'cancelled': ['Branch Code', 'Business Date', 'Invoice Number', 'Cancel/Void Date', 'Reason For Cancel/Void', 'Total'],
    'menu': ['SKU', 'ITEM NAME', 'ACCOUNT', 'STATUS'],
}
REPORT_NAMES = {
    'sales': '"Sales Items By Transaction" (Reports > Sales > Sales Items)',
    'cancelled': '"Cancelled Transactions" (Reports > Sales > Sale Transactions)',
    'menu': '"Export Items" (Business Settings > Items > three dots)',
}
PRICE_COL = re.compile(r'^Price-(.+?):(.+)$')
DATE_RE = re.compile(r'^\d{4}-\d{2}-\d{2}$')


async def ensure_indexes():
    await db.rista_sales_lines.create_index([('business_date', 1), ('branch_code', 1)])
    await db.rista_sales_lines.create_index('sku')
    await db.rista_cancelled.create_index([('business_date', 1), ('branch_code', 1)])
    await db.rista_menu.create_index('sku', unique=True)


# ---------------------------------------------------------------- reading files

def _text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace('\u2060', '').strip()


def _number(v):
    t = _text(v).replace(',', '')
    if not t:
        return 0.0
    try:
        return float(t)
    except ValueError:
        return 0.0


def _date_text(v):
    """Rista writes dates as text (CSV) or real dates (Excel). Keep YYYY-MM-DD[ HH:MM:SS] text either way."""
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d %H:%M:%S') if getattr(v, 'hour', 0) or getattr(v, 'minute', 0) else v.strftime('%Y-%m-%d')
    return _text(v)


def read_table(filename: str, content: bytes):
    """Return (headers, rows as dicts) from a Rista CSV or Excel download."""
    name = (filename or '').lower()
    if name.endswith('.xlsx') or content[:2] == b'PK':
        try:
            wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(400, 'This Excel file could not be read. Download the report again from Rista.')
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = [_text(h) for h in next(it, [])]
        rows = []
        for r in it:
            if r is None or all(c is None or _text(c) == '' for c in r):
                continue
            rows.append({h: (_date_text(c) if 'Date' in h else c) for h, c in zip(headers, r) if h})
        return headers, rows
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = content.decode('latin-1')
    text = text.replace('\u2060', '')
    reader = csv.DictReader(io.StringIO(text))
    headers = [h.strip() for h in (reader.fieldnames or [])]
    reader.fieldnames = headers
    rows = [r for r in reader if any(_text(v) for v in r.values() if not isinstance(v, list))]
    return headers, rows


def _check_kind(kind, headers):
    missing = [c for c in REQUIRED[kind] if c not in headers]
    if missing:
        raise HTTPException(400, f'This is not the Rista {REPORT_NAMES[kind]} file: columns missing: {", ".join(missing)}.')


def _pick(row, columns):
    out = {}
    for field, col, kind in columns:
        v = row.get(col)
        out[field] = _number(v) if kind == 'n' else _text(v)
    return out


def _day(v):
    return _text(v)[:10]


async def _read_upload(kind: str, file: UploadFile):
    content = await file.read()
    if not content:
        raise HTTPException(400, 'The file is empty.')
    if len(content) > MAX_BYTES:
        raise HTTPException(400, 'The file is larger than 40 MB. Download a shorter date range.')
    headers, rows = read_table(file.filename, content)
    _check_kind(kind, headers)
    if not rows:
        raise HTTPException(400, 'The file has no rows.')
    return headers, rows


async def _log_import(user, kind, filename, summary):
    doc = {'id': uid(), 'kind': kind, 'filename': filename, 'imported_at': now_iso(),
           'imported_by': user['name'], **summary}
    await db.rista_imports.insert_one(dict(doc))
    await audit(user, f'Imported Rista {kind} file', f"{filename}: {summary.get('rows', 0)} rows", doc['id'])
    return doc


async def _replace_by_branch_day(coll, docs, batch):
    """Store docs, then remove the older rows of the same branches and days.

    The new rows go in first, so if loading fails part way the new rows are removed again and what was loaded
    before stays untouched. Loading the same file again always repairs a failure that happened later."""
    days_by_branch = {}
    for d in docs:
        days_by_branch.setdefault(d['branch_code'], set()).add(d['business_date'])
    try:
        for i in range(0, len(docs), 2000):
            await db[coll].insert_many([dict(d) for d in docs[i:i + 2000]])
    except Exception:
        await db[coll].delete_many({'import_id': batch})
        raise HTTPException(500, 'Loading the file failed part way. Nothing was changed. Please try again.')
    replaced = 0
    try:
        for branch, days in days_by_branch.items():
            res = await db[coll].delete_many({'branch_code': branch, 'business_date': {'$in': sorted(days)},
                                              'import_id': {'$ne': batch}})
            replaced += res.deleted_count
    except Exception:
        raise HTTPException(500, 'Loading the file failed while replacing older rows. Please load the same file again.')
    return replaced


def _range(docs):
    days = sorted({d['business_date'] for d in docs})
    return {'date_from': days[0], 'date_to': days[-1], 'days': len(days),
            'branches': sorted({d['branch_name'] for d in docs})}


# ---------------------------------------------------------------- imports

@router.post('/import/sales')
async def import_sales(file: UploadFile = File(...), user=Depends(require_roles('admin'))):
    _, rows = await _read_upload('sales', file)
    batch = uid()
    docs = []
    skipped = 0
    for r in rows:
        d = _pick(r, SALES_COLUMNS)
        d['business_date'] = _day(d['business_date'])
        if not d['branch_code'] or not DATE_RE.match(d['business_date']) or not d['invoice_number']:
            skipped += 1
            continue
        d['sku'] = d['sku'].upper()
        d['import_id'] = batch
        docs.append(d)
    if not docs:
        raise HTTPException(400, 'No usable sales lines found (each line needs a branch code, business date and invoice number).')
    replaced = await _replace_by_branch_day('rista_sales_lines', docs, batch)
    _invalidate()
    orders = len({(d['branch_code'], d['invoice_number']) for d in docs})
    summary = {'rows': len(docs), 'skipped': skipped, 'replaced': replaced, 'orders': orders,
               'net': round(sum(d['net'] for d in docs), 2), **_range(docs)}
    return await _log_import(user, 'sales', file.filename, summary)


@router.post('/import/cancelled')
async def import_cancelled(file: UploadFile = File(...), user=Depends(require_roles('admin'))):
    _, rows = await _read_upload('cancelled', file)
    batch = uid()
    docs = []
    skipped = 0
    for r in rows:
        d = _pick(r, CANCELLED_COLUMNS)
        d['business_date'] = _day(d['business_date'])
        if not d['branch_code'] or not DATE_RE.match(d['business_date']):
            skipped += 1
            continue
        d['cancelled_by_platform'] = d['cancelled_by'].upper() == 'SYSTEM'
        d['import_id'] = batch
        docs.append(d)
    if not docs:
        raise HTTPException(400, 'No usable cancelled orders found.')
    replaced = await _replace_by_branch_day('rista_cancelled', docs, batch)
    _invalidate()
    summary = {'rows': len(docs), 'skipped': skipped, 'replaced': replaced,
               'total': round(sum(d['total'] for d in docs), 2), **_range(docs)}
    return await _log_import(user, 'cancelled', file.filename, summary)


@router.post('/import/menu')
async def import_menu(file: UploadFile = File(...), user=Depends(require_roles('admin'))):
    headers, rows = await _read_upload('menu', file)
    price_cols = [(h, *PRICE_COL.match(h).groups()) for h in headers if PRICE_COL.match(h)]
    batch = uid()
    docs = {}
    for r in rows:
        d = _pick(r, MENU_COLUMNS)
        d['sku'] = d['sku'].upper()
        if not d['sku'] or not d['name']:
            continue
        prices = []
        for col, branch, channel in price_cols:
            t = _text(r.get(col))
            if t:
                prices.append({'branch': branch.strip(), 'channel': channel.strip(), 'price': _number(t)})
        d['prices'] = prices
        values = [p['price'] for p in prices]
        d['price_min'] = min(values) if values else None
        d['price_max'] = max(values) if values else None
        d['branches'] = sorted({p['branch'] for p in prices})
        d['active'] = d['status'].lower() == 'active'
        d['import_id'] = batch
        docs[d['sku']] = d
    if not docs:
        raise HTTPException(400, 'No dishes found in the menu file.')
    # Replace dish by dish, then drop dishes that are no longer in the file: the menu is never empty half way.
    try:
        await db.rista_menu.bulk_write([ReplaceOne({'sku': d['sku']}, dict(d), upsert=True) for d in docs.values()], ordered=False)
        await db.rista_menu.delete_many({'import_id': {'$ne': batch}})
    except Exception:
        raise HTTPException(500, 'Loading the menu failed part way. Please load the same file again.')
    _invalidate()
    summary = {'rows': len(docs), 'active': sum(1 for d in docs.values() if d['active']),
               'accounts': len({d['account'] for d in docs.values()}), 'price_columns': len(price_cols)}
    return await _log_import(user, 'menu', file.filename, summary)


@router.get('/imports')
async def list_imports(user=Depends(require_roles(*VIEW_ROLES))):
    return clean_list(await db.rista_imports.find({}).sort('imported_at', -1).to_list(200))


# ---------------------------------------------------------------- sales report

BRAND_FIELDS = {'business_brand', 'account_name'}


def _filters(date_from='', date_to='', branch='', brand='', brand_field='business_brand', channel='', source=''):
    if brand_field not in BRAND_FIELDS:
        raise HTTPException(400, 'brand_field must be business_brand or account_name')
    q = {}
    if date_from or date_to:
        q['business_date'] = {}
        if date_from:
            q['business_date']['$gte'] = date_from
        if date_to:
            q['business_date']['$lte'] = date_to
    if branch:
        q['branch_code'] = branch
    if brand:
        q[brand_field] = brand
    if channel:
        q['channel'] = channel
    if source:
        q['order_source'] = source
    return q


def _round(rows, *keys):
    for r in rows:
        for k in keys:
            if r.get(k) is not None:
                r[k] = round(r[k], 2)
    return rows


async def _group(match, key, extra_sort=None, limit=None):
    pipeline = [
        {'$match': match},
        {'$group': {'_id': {'k': key, 'branch': '$branch_code', 'inv': '$invoice_number'},
                    'net': {'$sum': '$net'}, 'gross': {'$sum': '$gross'}, 'discounts': {'$sum': '$discounts'}}},
        {'$group': {'_id': '$_id.k', 'orders': {'$sum': 1}, 'net': {'$sum': '$net'},
                    'gross': {'$sum': '$gross'}, 'discounts': {'$sum': '$discounts'}}},
        {'$sort': extra_sort or {'net': -1}},
    ]
    if limit:
        pipeline.append({'$limit': limit})
    rows = await db.rista_sales_lines.aggregate(pipeline).to_list(5000)
    return _round([{'key': r['_id'] or '(blank)', 'orders': r['orders'], 'net': r['net'], 'gross': r['gross'],
                    'discounts': r['discounts']} for r in rows], 'net', 'gross', 'discounts')


@router.get('/sales/filters')
async def sales_filters(user=Depends(require_roles(*VIEW_ROLES))):
    return await _cached(('filters',), _filters_data)


async def _filters_data():
    coll = db.rista_sales_lines
    branches = await coll.aggregate([{'$group': {'_id': {'code': '$branch_code', 'name': '$branch_name'}}}]).to_list(100)
    first = await coll.find({}, {'business_date': 1}).sort('business_date', 1).limit(1).to_list(1)
    last = await coll.find({}, {'business_date': 1}).sort('business_date', -1).limit(1).to_list(1)
    return {
        'branches': sorted([{'code': b['_id']['code'], 'name': b['_id']['name']} for b in branches], key=lambda b: b['name']),
        'business_brands': sorted(x for x in await coll.distinct('business_brand') if x),
        'account_names': sorted(x for x in await coll.distinct('account_name') if x),
        'channels': sorted(x for x in await coll.distinct('channel') if x),
        'sources': sorted(x for x in await coll.distinct('order_source') if x),
        'date_from': first[0]['business_date'] if first else None,
        'date_to': last[0]['business_date'] if last else None,
    }


@router.get('/sales/summary')
async def sales_summary(date_from: str = '', date_to: str = '', branch: str = '', brand: str = '',
                        brand_field: str = 'business_brand', channel: str = '', source: str = '',
                        user=Depends(require_roles(*VIEW_ROLES))):
    match = _filters(date_from, date_to, branch, brand, brand_field, channel, source)
    key = ('summary', date_from, date_to, branch, brand, brand_field, channel, source)
    return await _cached(key, lambda: _summary(match, brand_field))


async def _summary(match, brand_field):
    totals = await db.rista_sales_lines.aggregate([
        {'$match': match},
        {'$group': {'_id': None, 'lines': {'$sum': 1}, 'gross': {'$sum': '$gross'}, 'discounts': {'$sum': '$discounts'},
                    'net': {'$sum': '$net'}, 'taxes': {'$sum': '$taxes'}, 'aggregator_tax': {'$sum': '$aggregator_tax'},
                    'charges': {'$sum': {'$add': ['$direct_charge', '$indirect_charge']}},
                    'qty': {'$sum': {'$cond': [{'$eq': ['$line_type', 'Item']}, '$qty', 0]}},
                    'orders': {'$addToSet': {'b': '$branch_code', 'i': '$invoice_number'}},
                    'days': {'$addToSet': '$business_date'}}},
        {'$project': {'_id': 0, 'lines': 1, 'gross': 1, 'discounts': 1, 'net': 1, 'taxes': 1, 'aggregator_tax': 1,
                      'charges': 1, 'qty': 1, 'orders': {'$size': '$orders'}, 'days': {'$size': '$days'}}},
    ]).to_list(1)
    t = totals[0] if totals else {'lines': 0, 'gross': 0, 'discounts': 0, 'net': 0, 'taxes': 0,
                                  'aggregator_tax': 0, 'charges': 0, 'qty': 0, 'orders': 0, 'days': 0}
    _round([t], 'gross', 'discounts', 'net', 'taxes', 'aggregator_tax', 'charges', 'qty')
    t['avg_order'] = round(t['net'] / t['orders'], 2) if t['orders'] else 0
    t['orders_per_day'] = round(t['orders'] / t['days'], 1) if t['days'] else 0

    def share(rows):
        for r in rows:
            r['share'] = round(r['net'] / t['net'] * 100, 1) if t['net'] else 0
        return rows

    dishes_q = db.rista_sales_lines.aggregate([
        {'$match': {**match, 'line_type': 'Item'}},
        {'$group': {'_id': '$sku', 'item_name': {'$last': '$item_name'}, 'account_name': {'$last': '$account_name'},
                    'qty': {'$sum': '$qty'}, 'net': {'$sum': '$net'}}},
        {'$sort': {'net': -1}}, {'$limit': 50},
    ]).to_list(50)
    by_day, by_brand, by_source, by_channel, by_branch, dishes = await asyncio.gather(
        _group(match, '$business_date', {'_id': 1}), _group(match, f'${brand_field}'),
        _group(match, '$order_source'), _group(match, '$channel'), _group(match, '$branch_name'), dishes_q)
    return {
        'totals': t,
        'by_day': by_day,
        'by_brand': share(by_brand),
        'by_source': share(by_source),
        'by_channel': share(by_channel),
        'by_branch': share(by_branch),
        'top_dishes': share(_round([{'sku': d['_id'], 'item_name': d['item_name'], 'account_name': d['account_name'],
                                     'qty': d['qty'], 'net': d['net']} for d in dishes], 'qty', 'net')),
    }


EXPORT_HEADERS = [
    ('business_date', 'Business Date'), ('branch_name', 'Branch'), ('business_brand', 'Business Brand'),
    ('account_name', 'Account Name'), ('order_source', 'Order Source'), ('channel', 'Channel'),
    ('invoice_number', 'Invoice Number'), ('invoice_date', 'Invoice Date'), ('source_order_number', 'Source Order Number'),
    ('sku', 'SKU'), ('line_type', 'Type'), ('item_name', 'Item Name'), ('variants', 'Variant(s)'),
    ('category', 'Category'), ('qty', 'Qty'), ('unit_rate', 'Unit Rate'), ('gross', 'Gross Amount'),
    ('discounts', 'Discounts'), ('net', 'Net Amount'), ('taxes', 'Taxes'),
    ('aggregator_tax', 'Tax Collected by Aggregator'), ('direct_charge', 'Direct Charge'),
    ('indirect_charge', 'Indirect Charge'), ('payment_modes', 'Payment Modes'),
]
CANCELLED_HEADERS = [
    ('business_date', 'Business Date'), ('branch_name', 'Branch'), ('business_brand', 'Business Brand'),
    ('order_source', 'Order Source'), ('channel', 'Channel'), ('invoice_number', 'Invoice Number'),
    ('source_order_number', 'Source Order Number'), ('order_date', 'Order Date'), ('cancel_date', 'Cancelled At'),
    ('time_to_cancel', 'Minutes To Cancel'), ('status', 'Status'), ('fulfillment_status', 'Stage When Cancelled'),
    ('cancelled_by', 'Cancelled By'), ('reason', 'Reason'), ('source_reason', 'Platform Reason'),
    ('gross', 'Gross Amount'), ('discounts', 'Discounts'), ('net', 'Net Amount'), ('total', 'Total'),
]


def _build_file(rows, headers, fmt):
    if fmt == 'csv':
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([h for _, h in headers])
        for r in rows:
            w.writerow([r.get(k, '') for k, _ in headers])
        return ('\ufeff' + buf.getvalue()).encode('utf-8'), 'text/csv'
    else:
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Data')
        ws.freeze_panes = 'A2'
        head = []
        for _, h in headers:
            c = WriteOnlyCell(ws, value=h)
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1E293B')
            head.append(c)
        ws.append(head)
        for r in rows:
            ws.append([r.get(k, '') for k, _ in headers])
        data = BytesIO()
        wb.save(data)
        return data.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


async def _file_response(rows, headers, name, fmt):
    # Building a month of lines takes a few seconds: do it off the event loop so other requests keep flowing.
    content, media = await asyncio.to_thread(_build_file, rows, headers, fmt)
    return Response(content, media_type=media, headers={'Content-Disposition': f'attachment; filename="{name}.{fmt}"'})


def _check_format(fmt):
    if fmt not in ('csv', 'xlsx'):
        raise HTTPException(400, 'format must be csv or xlsx')


@router.get('/sales/export')
async def export_sales(format: str = 'xlsx', date_from: str = '', date_to: str = '', branch: str = '', brand: str = '',
                       brand_field: str = 'business_brand', channel: str = '', source: str = '',
                       user=Depends(require_roles(*VIEW_ROLES))):
    _check_format(format)
    match = _filters(date_from, date_to, branch, brand, brand_field, channel, source)
    count = await db.rista_sales_lines.count_documents(match)
    if count > EXPORT_MAX_LINES:
        raise HTTPException(400, f'This download would have {count:,} lines, more than the {EXPORT_MAX_LINES:,} allowed at once. '
                                 'Choose a shorter date range, a branch or a brand.')
    rows = await db.rista_sales_lines.find(match, {'_id': 0}).sort([('business_date', 1), ('invoice_date', 1)]).to_list(EXPORT_MAX_LINES)
    await audit(user, 'Downloaded Rista sales', f'{len(rows)} lines {date_from}..{date_to} {brand}')
    return await _file_response(rows, EXPORT_HEADERS, f"rista-sales-{date_from or 'all'}-{date_to or 'all'}", format)


# ---------------------------------------------------------------- cancelled orders

@router.get('/cancelled')
async def list_cancelled(date_from: str = '', date_to: str = '', branch: str = '', brand: str = '',
                         channel: str = '', source: str = '', user=Depends(require_roles(*VIEW_ROLES))):
    match = _filters(date_from, date_to, branch, brand, 'business_brand', channel, source)
    key = ('cancelled', date_from, date_to, branch, brand, channel, source)
    return await _cached(key, lambda: _cancelled_data(match))


async def _cancelled_data(match):
    rows = await db.rista_cancelled.find(match, {'_id': 0, 'import_id': 0}).sort('cancel_date', -1).to_list(20000)
    by_reason = {}
    for r in rows:
        k = r['reason'] or '(blank)'
        by_reason[k] = by_reason.get(k, 0) + 1
    return {
        'rows': rows,
        'totals': {'orders': len(rows), 'total': round(sum(r['total'] for r in rows), 2),
                   'by_platform': sum(1 for r in rows if r['cancelled_by_platform']),
                   'by_staff': sum(1 for r in rows if not r['cancelled_by_platform'])},
        'by_reason': sorted([{'reason': k, 'orders': v} for k, v in by_reason.items()], key=lambda x: -x['orders']),
    }


@router.get('/cancelled/export')
async def export_cancelled(format: str = 'xlsx', date_from: str = '', date_to: str = '', branch: str = '',
                           brand: str = '', channel: str = '', source: str = '',
                           user=Depends(require_roles(*VIEW_ROLES))):
    _check_format(format)
    match = _filters(date_from, date_to, branch, brand, 'business_brand', channel, source)
    rows = await db.rista_cancelled.find(match, {'_id': 0}).sort('cancel_date', 1).to_list(20000)
    return await _file_response(rows, CANCELLED_HEADERS, f"rista-cancelled-{date_from or 'all'}-{date_to or 'all'}", format)


# ---------------------------------------------------------------- menu

MENU_EXPORT_HEADERS = [
    ('sku', 'SKU'), ('name', 'Dish'), ('variant', 'Variant'), ('account', 'Brand Account'), ('category', 'Category'),
    ('status', 'Status'), ('price_min', 'Lowest Price'), ('price_max', 'Highest Price'), ('sold_qty', 'Sold Qty'),
    ('sold_net', 'Sold Net Amount'), ('has_recipe_text', 'Recipe In This App'),
]


def _menu_match(q, account, status):
    match = {}
    if q.strip():
        rx = {'$regex': re.escape(q.strip()), '$options': 'i'}
        match['$or'] = [{'name': rx}, {'sku': rx}, {'category': rx}]
    if account:
        match['account'] = account
    if status == 'active':
        match['active'] = True
    elif status == 'inactive':
        match['active'] = False
    return match


async def _sold_by_sku():
    async def compute():
        return await db.rista_sales_lines.aggregate([
            {'$match': {'line_type': 'Item'}},
            {'$group': {'_id': '$sku', 'item_name': {'$last': '$item_name'}, 'account_name': {'$last': '$account_name'},
                        'lines': {'$sum': 1}, 'qty': {'$sum': '$qty'}, 'net': {'$sum': '$net'}}},
            {'$sort': {'net': -1}}]).to_list(10000)
    return await _cached(('sold',), compute)


async def _recipe_codes():
    return {r['code'] for r in await db.recipes.find({}, {'code': 1}).to_list(10000) if r.get('code')}


async def _menu_rows(q, account, status):
    """Every matching dish with how much of it sold and whether a recipe here is linked to it."""
    async def compute():
        # Only what the list shows; full details and prices come from /menu/{sku}.
        fields = {'_id': 0, 'sku': 1, 'name': 1, 'variant': 1, 'account': 1, 'category': 1, 'status': 1, 'active': 1,
                  'price_min': 1, 'price_max': 1}
        return await db.rista_menu.find(_menu_match(q, account, status), fields).sort([('account', 1), ('name', 1)]).to_list(5000)
    base = await _cached(('menu', q.strip().lower(), account, status), compute)
    sold = {s['_id']: s for s in await _sold_by_sku()}
    codes = await _recipe_codes()  # read fresh: recipes change outside Rista file loads
    out = []
    for r in base:
        s = sold.get(r['sku'])
        out.append({**r, 'sold_qty': s['qty'] if s else 0, 'sold_net': round(s['net'], 2) if s else 0,
                    'has_recipe': r['sku'] in codes})
    return out


@router.get('/menu')
async def list_menu(q: str = '', account: str = '', status: str = '', page: int = 1, page_size: int = 50,
                    user=Depends(require_roles(*VIEW_ROLES))):
    rows = await _menu_rows(q, account, status)
    page_size = min(max(page_size, 1), 200)
    pages = max(1, math.ceil(len(rows) / page_size))
    page = min(max(page, 1), pages)
    return {'rows': rows[(page - 1) * page_size:page * page_size], 'total': len(rows), 'page': page,
            'pages': pages, 'page_size': page_size}


@router.get('/menu/export')
async def export_menu(format: str = 'xlsx', q: str = '', account: str = '', status: str = '',
                      user=Depends(require_roles(*VIEW_ROLES))):
    _check_format(format)
    rows = [{**r, 'has_recipe_text': 'Yes' if r['has_recipe'] else 'No'} for r in await _menu_rows(q, account, status)]
    return await _file_response(rows, MENU_EXPORT_HEADERS, 'rista-menu', format)


@router.get('/menu/accounts')
async def menu_accounts(user=Depends(require_roles(*VIEW_ROLES))):
    async def compute():
        return sorted(x for x in await db.rista_menu.distinct('account') if x)
    return await _cached(('accounts',), compute)


@router.get('/menu/checks')
async def menu_checks(user=Depends(require_roles(*VIEW_ROLES))):
    """Dishes sold that are missing from the menu file, and dishes sold that have no recipe in our app."""
    async def menu_skus():
        return set(await db.rista_menu.distinct('sku'))
    skus = await _cached(('menu_skus',), menu_skus)
    codes = await _recipe_codes()
    rows = [{'sku': s['_id'], 'item_name': s['item_name'], 'account_name': s['account_name'], 'lines': s['lines'],
             'qty': s['qty'], 'net': round(s['net'], 2)} for s in await _sold_by_sku()]
    return {
        'menu_loaded': bool(skus),
        'sold_dishes': len(rows),
        'not_in_menu': [r for r in rows if r['sku'] not in skus] if skus else [],
        'without_recipe': [r for r in rows if r['sku'] not in codes],
    }


@router.get('/menu/{sku}')
async def menu_item(sku: str, user=Depends(require_roles(*VIEW_ROLES))):
    d = await db.rista_menu.find_one({'sku': sku.upper()}, {'_id': 0, 'import_id': 0})
    if not d:
        raise HTTPException(404, 'Dish not found in the Rista menu')
    return d
