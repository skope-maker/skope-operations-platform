from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from db import db, today_str
from auth import require_roles
from inventory import to_base, post_txn
from routers_recipes import SubRecipeIn, SubRecipeItemIn, RecipeIn, RecipeItemIn, create_subrecipe, create_recipe
from routers_sales import SaleIn, SaleItemIn, create_sale
from routers_masters import ItemIn, create_item

router = APIRouter(prefix='/bulk', tags=['bulk'])

HEADER_FILL = PatternFill('solid', fgColor='1E293B')
MANDATORY_FILL = PatternFill('solid', fgColor='0F766E')

SPECS = {
    'item': {
        'sheet': 'Items',
        'label': 'Item',
        'columns': [
            ('Item Code', True, 'Unique code for the item, e.g. ITM-ONION.'),
            ('Item Name', True, 'Display name, e.g. Onion.'),
            ('Category', False, 'Category master name. Leave blank if none.'),
            ('Subcategory', False, 'Free text subcategory.'),
            ('Purchase Unit', True, 'Unit master code you buy in, e.g. KG, LTR, BOX.'),
            ('Consumption Unit', True, 'Unit master code you consume/track in, e.g. GM, ML. Must be the same family as purchase unit.'),
            ('Conversion Factor', False, 'Purchase unit to consumption unit factor. Blank = 1.'),
            ('Min Stock', False, 'Minimum stock in consumption unit. Blank = 0.'),
            ('Reorder Level', False, 'Reorder trigger level in consumption unit. Blank = 0.'),
            ('Max Stock', False, 'Maximum / par stock in consumption unit. Blank = 0.'),
            ('Rate', False, 'Purchase rate per purchase unit (number). Blank = 0.'),
            ('Supplier', False, 'Supplier master name. Leave blank if none.'),
            ('Opening Qty', False, 'Optional opening stock quantity (number > 0). Requires Opening Unit and Opening Location.'),
            ('Opening Unit', False, 'Unit master code of the opening quantity. Same family as consumption unit.'),
            ('Opening Location', False, 'Location/kitchen master code where the opening stock sits, e.g. STORE.'),
        ],
        'examples': [
            ['ITM-ONION', 'Onion', 'Vegetables', 'Fresh', 'KG', 'GM', 1000, 5000, 10000, 50000, 40, 'Local Mandi', 25, 'KG', 'STORE'],
            ['ITM-OIL', 'Refined Oil', 'Groceries', '', 'LTR', 'ML', 1000, 2000, 5000, 20000, 120, 'Local Mandi', '', '', ''],
            ['ITM-CHICKEN', 'Chicken', 'Meat', 'Fresh', 'KG', 'GM', 1000, 3000, 6000, 20000, 220, '', 10, 'KG', 'STORE'],
        ],
    },
    'subrecipe': {
        'sheet': 'SubRecipes',
        'label': 'Sub-Recipe',
        'columns': [
            ('Sub Recipe Code', True, 'Unique code for the sub-recipe. Repeat the same code on every ingredient row of that sub-recipe.'),
            ('Sub Recipe Name', True, 'Display name, e.g. Chicken Gravy Base.'),
            ('Brand Code', False, 'Brand master code. Leave blank for shared sub-recipes.'),
            ('Output Item Code', True, 'Item master code of the prepared output item that this batch produces.'),
            ('Yield Qty', True, 'Final produced quantity of one batch (number > 0). Same on every row of the sub-recipe.'),
            ('Yield Unit', True, 'Unit master code of the yield, e.g. KG, GM, LTR, ML, PCS.'),
            ('Ingredient Item Code', True, 'Item master code of the raw ingredient consumed.'),
            ('Ingredient Qty', True, 'Quantity of the ingredient for one batch (number > 0).'),
            ('Ingredient Unit', True, 'Unit master code of the ingredient quantity.'),
        ],
        'examples': [
            ['SR-GRAVY', 'Chicken Gravy Base', '', 'ITM-GRAVY', 10, 'KG', 'ITM-ONION', 6, 'KG'],
            ['SR-GRAVY', 'Chicken Gravy Base', '', 'ITM-GRAVY', 10, 'KG', 'ITM-TOMATO', 4, 'KG'],
            ['SR-GRAVY', 'Chicken Gravy Base', '', 'ITM-GRAVY', 10, 'KG', 'ITM-OIL', 500, 'ML'],
        ],
    },
    'recipe': {
        'sheet': 'Recipes',
        'label': 'Main Recipe',
        'columns': [
            ('Recipe Code', True, 'Unique code for the recipe. Repeat on every ingredient row of that recipe.'),
            ('Recipe Name', True, 'Menu item name, e.g. Butter Chicken.'),
            ('Brand Code', True, 'Brand master code the recipe belongs to.'),
            ('Category', False, 'Free text menu category, e.g. Main Course.'),
            ('Portion Size', False, 'Free text, e.g. 250 GM.'),
            ('Selling Unit', False, 'Defaults to Portion.'),
            ('Selling Price', False, 'Menu price per portion (number).'),
            ('Effective Date', False, 'YYYY-MM-DD. Blank = effective immediately.'),
            ('Line Type', True, 'ITEM for a raw item line, SUBRECIPE for a sub-recipe line.'),
            ('Item / Sub Recipe Code', True, 'Item master code when Line Type = ITEM, sub-recipe code when Line Type = SUBRECIPE.'),
            ('Qty Per Portion', True, 'Quantity used for one portion (number > 0).'),
            ('Unit', True, 'Unit master code of the quantity.'),
        ],
        'examples': [
            ['RCP-BC', 'Butter Chicken', 'BRAND1', 'Main Course', '250 GM', 'Portion', 380, '2026-06-01', 'SUBRECIPE', 'SR-GRAVY', 200, 'GM'],
            ['RCP-BC', 'Butter Chicken', 'BRAND1', 'Main Course', '250 GM', 'Portion', 380, '2026-06-01', 'ITEM', 'ITM-CHICKEN', 180, 'GM'],
            ['RCP-BC', 'Butter Chicken', 'BRAND1', 'Main Course', '250 GM', 'Portion', 380, '2026-06-01', 'ITEM', 'ITM-BUTTER', 20, 'GM'],
        ],
    },
    'sales': {
        'sheet': 'Sales',
        'label': 'Daily Sales',
        'columns': [
            ('Date', True, 'Sales date in YYYY-MM-DD format.'),
            ('Brand Code', True, 'Brand master code.'),
            ('Kitchen Code', True, 'Location/kitchen master code where the sale happened.'),
            ('Recipe Code', True, 'Recipe master code of the product sold.'),
            ('Qty Sold', True, 'Total portions sold (number >= 0).'),
            ('Cancelled', False, 'Cancelled portions. Cannot exceed Qty Sold. Blank = 0.'),
            ('Complimentary', False, 'Complimentary portions. Cannot exceed Qty Sold - Cancelled. Blank = 0.'),
        ],
        'examples': [
            ['2026-06-01', 'BRAND1', 'NORMK', 'RCP-BC', 42, 2, 1],
            ['2026-06-01', 'BRAND1', 'NORMK', 'RCP-DAL', 18, 0, 0],
            ['2026-06-02', 'BRAND1', 'NORMK', 'RCP-BC', 37, 0, 0],
        ],
    },
}


def _instructions_sheet(wb, kind):
    spec = SPECS[kind]
    ws = wb.create_sheet('Instructions')
    ws['A1'] = f"{spec['label']} Bulk Upload — Instructions"
    ws['A1'].font = Font(bold=True, size=14)
    rows = [
        [],
        ['Column', 'Mandatory', 'Meaning / Accepted values'],
    ]
    for name, mandatory, desc in spec['columns']:
        rows.append([name, 'YES' if mandatory else 'Optional', desc])
    rows += [
        [],
        ['How to upload'],
        ['1. Fill the data in the first sheet only. Do not rename or reorder the columns.'],
        ['2. One row per ingredient / product line. Repeat the header columns on each row of the same group.'],
        ['3. Upload the file in the app, review the preview, fix any red rows, then confirm the import.'],
        [],
        ['Common validation errors'],
        ['Mandatory field is blank'],
        ['Item / sub-recipe / recipe / brand / kitchen code does not exist in masters'],
        ['Unit code not present in the Units master, or not convertible with the item unit'],
        ['Quantity is zero, negative or not a number'],
        ['Code already exists in the system (duplicate record)'],
        ['Same code repeated with different header values inside the file'],
    ]
    for r in rows:
        ws.append(r)
    for c in ('A', 'B', 'C'):
        ws.column_dimensions[c].width = {'A': 30, 'B': 14, 'C': 95}[c]
    for cell in ws[3]:
        if cell.value:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = HEADER_FILL
    for row in ws.iter_rows(min_row=4):
        row[2].alignment = Alignment(wrap_text=True, vertical='top')
    return ws


def build_template(kind: str) -> BytesIO:
    spec = SPECS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = spec['sheet']
    headers = [f"{n}*" if m else n for n, m, _ in spec['columns']]
    ws.append(headers)
    for i, (name, mandatory, _) in enumerate(spec['columns'], start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = MANDATORY_FILL if mandatory else HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = max(16, len(name) + 6)
    for ex in spec['examples']:
        ws.append(ex)
    ws.freeze_panes = 'A2'
    note = ws.cell(row=len(spec['examples']) + 3, column=1)
    note.value = 'Columns marked * are mandatory. Delete these example rows before uploading. See the Instructions sheet.'
    note.font = Font(italic=True, color='B91C1C')
    _instructions_sheet(wb, kind)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get('/template/{kind}')
async def download_template(kind: str, user=Depends(require_roles('admin', 'chef', 'store_manager'))):
    if kind not in SPECS:
        raise HTTPException(404, 'Unknown template')
    buf = build_template(kind)
    filename = f"{kind}-bulk-upload-template.xlsx"
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             headers={'Content-Disposition': f'attachment; filename="{filename}"'})


# ---------------- parsing helpers ----------------
def _norm(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _num(v):
    if v is None or _norm(v) == '':
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _date(v):
    s = _norm(v)
    if not s:
        return ''
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    return s[:10]


async def _read_rows(kind: str, file: UploadFile):
    spec = SPECS[kind]
    data = await file.read()
    if not data:
        raise HTTPException(400, 'Uploaded file is empty')
    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except Exception:
        raise HTTPException(400, 'Could not read the file. Upload the .xlsx template filled with your data.')
    ws = wb[spec['sheet']] if spec['sheet'] in wb.sheetnames else wb[
        next((n for n in wb.sheetnames if n != 'Instructions'), wb.sheetnames[0])]
    ncols = len(spec['columns'])
    rows = []
    for idx, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = list(raw)[:ncols] + [None] * max(0, ncols - len(raw or []))
        if all(_norm(v) == '' for v in values):
            continue
        first = _norm(values[0])
        if first.lower().startswith('columns marked'):
            continue
        rows.append((idx, values))
    if not rows:
        raise HTTPException(400, 'No data rows found in the file')
    return rows


async def _masters():
    return {
        'items': {i['code'].upper(): i for i in await db.items.find({'deleted': {'$ne': True}}, {'_id': 0}).to_list(10000)},
        'units': {u['code'].upper(): u for u in await db.units.find({'deleted': {'$ne': True}}, {'_id': 0}).to_list(500)},
        'brands': {b['code'].upper(): b for b in await db.brands.find({'deleted': {'$ne': True}}, {'_id': 0}).to_list(500)},
        'kitchens': {k['code'].upper(): k for k in await db.kitchens.find({'deleted': {'$ne': True}}, {'_id': 0}).to_list(500)},
        'subs': {s['code'].upper(): s for s in await db.sub_recipes.find({}, {'_id': 0}).to_list(2000)},
        'recipes': {r['code'].upper(): r for r in await db.recipes.find({}, {'_id': 0}).to_list(5000)},
        'categories': {c['name'].upper(): c for c in await db.categories.find({'deleted': {'$ne': True}}, {'_id': 0}).to_list(2000)},
        'suppliers': {s['name'].upper(): s for s in await db.suppliers.find({'deleted': {'$ne': True}}, {'_id': 0}).to_list(2000)},
    }


def _err(errors, row, field, reason, value):
    errors.append({'row': row, 'field': field, 'reason': reason, 'value': _norm(value)})


# ---------------- validators ----------------
async def validate_item(rows, m):
    groups, errors, seen = [], [], set()
    for row, v in rows:
        (code, name, cat, subcat, pu, cu, conv, mins, reorder, maxs, rate, supplier, oqty, ounit, oloc) = v
        code_u = _norm(code).upper()
        e = []
        if not code_u:
            _err(e, row, 'Item Code', 'Mandatory field is blank', code)
        elif code_u in seen:
            _err(e, row, 'Item Code', 'Item code repeated within the file', code)
        elif code_u in m['items']:
            _err(e, row, 'Item Code', 'Item code already exists in the system', code)
        if code_u:
            seen.add(code_u)
        if not _norm(name):
            _err(e, row, 'Item Name', 'Mandatory field is blank', name)
        pu_u, cu_u = _norm(pu).upper(), _norm(cu).upper()
        pud, cud = m['units'].get(pu_u), m['units'].get(cu_u)
        if not pu_u:
            _err(e, row, 'Purchase Unit', 'Mandatory field is blank', pu)
        elif not pud:
            _err(e, row, 'Purchase Unit', 'Unit code not found in Unit master', pu)
        if not cu_u:
            _err(e, row, 'Consumption Unit', 'Mandatory field is blank', cu)
        elif not cud:
            _err(e, row, 'Consumption Unit', 'Unit code not found in Unit master', cu)
        if pud and cud and pud['family'] != cud['family']:
            _err(e, row, 'Purchase / Consumption Unit', 'Purchase and consumption units are not convertible (different families)', f'{pu_u}/{cu_u}')
        cat_u = _norm(cat).upper()
        cat_doc = m['categories'].get(cat_u) if cat_u else None
        if cat_u and not cat_doc:
            _err(e, row, 'Category', 'Category not found in Category master', cat)
        sup_u = _norm(supplier).upper()
        sup_doc = m['suppliers'].get(sup_u) if sup_u else None
        if sup_u and not sup_doc:
            _err(e, row, 'Supplier', 'Supplier not found in Supplier master', supplier)
        conv_v = _num(conv)
        if _norm(conv) and (conv_v is None or conv_v <= 0):
            _err(e, row, 'Conversion Factor', 'Must be a number greater than zero', conv)
        for fname, fval in (('Min Stock', mins), ('Reorder Level', reorder), ('Max Stock', maxs), ('Rate', rate)):
            n = _num(fval)
            if _norm(fval) and (n is None or n < 0):
                _err(e, row, fname, 'Must be a number of zero or more', fval)
        opening = None
        if any(_norm(x) for x in (oqty, ounit, oloc)):
            oq = _num(oqty)
            ou = _norm(ounit).upper()
            od = m['kitchens'].get(_norm(oloc).upper())
            if oq is None or oq <= 0:
                _err(e, row, 'Opening Qty', 'Must be a number greater than zero when opening stock is given', oqty)
            if not ou:
                _err(e, row, 'Opening Unit', 'Required when opening stock is given', ounit)
            elif ou not in m['units']:
                _err(e, row, 'Opening Unit', 'Unit code not found in Unit master', ounit)
            elif cud and m['units'][ou]['family'] != cud['family']:
                _err(e, row, 'Opening Unit', 'Opening unit is not convertible with the consumption unit', ounit)
            if not _norm(oloc):
                _err(e, row, 'Opening Location', 'Required when opening stock is given', oloc)
            elif not od:
                _err(e, row, 'Opening Location', 'Location code not found in Location master', oloc)
            if oq and oq > 0 and ou in m['units'] and od:
                opening = {'qty': oq, 'unit': ou, 'location_id': od['id'], 'location_name': od['name']}
        groups.append({
            'key': code_u or f'__row{row}', 'code': code_u, 'name': _norm(name),
            'category_id': (cat_doc or {}).get('id', ''), 'subcategory': _norm(subcat),
            'purchase_unit': pu_u, 'consumption_unit': cu_u,
            'conversion_factor': conv_v or 1, 'min_stock': _num(mins) or 0,
            'reorder_level': _num(reorder) or 0, 'max_stock': _num(maxs) or 0,
            'rate': _num(rate) or 0, 'supplier_id': (sup_doc or {}).get('id', ''),
            'opening': opening, 'rows': [row], 'items': [], 'errors': e, 'valid': not e,
        })
        errors += e
    return groups, errors


async def validate_subrecipe(rows, m):
    groups, order, errors = {}, [], []
    for row, v in rows:
        code, name, brand, out_code, yqty, yunit, ing_code, ing_qty, ing_unit = v
        code_u = _norm(code).upper()
        row_errors = []
        if not code_u:
            _err(row_errors, row, 'Sub Recipe Code', 'Mandatory field is blank', code)
        if not _norm(name):
            _err(row_errors, row, 'Sub Recipe Name', 'Mandatory field is blank', name)
        out = m['items'].get(_norm(out_code).upper())
        if not _norm(out_code):
            _err(row_errors, row, 'Output Item Code', 'Mandatory field is blank', out_code)
        elif not out:
            _err(row_errors, row, 'Output Item Code', 'Item code not found in Item master', out_code)
        yq = _num(yqty)
        if yq is None or yq <= 0:
            _err(row_errors, row, 'Yield Qty', 'Must be a number greater than zero', yqty)
        if _norm(yunit).upper() not in m['units']:
            _err(row_errors, row, 'Yield Unit', 'Unit code not found in Unit master', yunit)
        ing = m['items'].get(_norm(ing_code).upper())
        if not _norm(ing_code):
            _err(row_errors, row, 'Ingredient Item Code', 'Mandatory field is blank', ing_code)
        elif not ing:
            _err(row_errors, row, 'Ingredient Item Code', 'Item code not found in Item master', ing_code)
        iq = _num(ing_qty)
        if iq is None or iq <= 0:
            _err(row_errors, row, 'Ingredient Qty', 'Must be a number greater than zero', ing_qty)
        if _norm(ing_unit).upper() not in m['units']:
            _err(row_errors, row, 'Ingredient Unit', 'Unit code not found in Unit master', ing_unit)
        brand_u = _norm(brand).upper()
        brand_doc = m['brands'].get(brand_u) if brand_u else None
        if brand_u and not brand_doc:
            _err(row_errors, row, 'Brand Code', 'Brand code not found in Brand master', brand)

        key = code_u or f"__row{row}"
        if key not in groups:
            if code_u and code_u in m['subs']:
                _err(row_errors, row, 'Sub Recipe Code', 'Sub-recipe code already exists in the system', code)
            groups[key] = {'key': key, 'code': code_u, 'name': _norm(name),
                           'brand_id': (brand_doc or {}).get('id', ''),
                           'output_item_id': (out or {}).get('id', ''),
                           'output_item_name': (out or {}).get('name', ''),
                           'yield_qty': yq or 0, 'yield_unit': _norm(yunit).upper(),
                           'rows': [], 'items': [], 'errors': []}
            order.append(key)
        g = groups[key]
        if g['rows'] and (g['name'] != _norm(name) or g['yield_qty'] != (yq or 0)):
            _err(row_errors, row, 'Sub Recipe Name / Yield Qty',
                 'Header values differ from the first row of this sub-recipe code', name)
        if ing and any(x['item_id'] == ing['id'] for x in g['items']):
            _err(row_errors, row, 'Ingredient Item Code', 'Same ingredient repeated within this sub-recipe', ing_code)
        g['rows'].append(row)
        if row_errors:
            g['errors'] += row_errors
            errors += row_errors
        elif ing:
            g['items'].append({'item_id': ing['id'], 'item_name': ing['name'], 'qty': iq,
                               'unit': _norm(ing_unit).upper(), 'row': row})
    for key in order:
        g = groups[key]
        if not g['items'] and not g['errors']:
            _err(g['errors'], g['rows'][0], 'Ingredient Item Code', 'Sub-recipe has no valid ingredient rows', '')
            errors += g['errors']
        g['valid'] = not g['errors']
    return [groups[k] for k in order], errors


async def validate_recipe(rows, m):
    groups, order, errors = {}, [], []
    for row, v in rows:
        (code, name, brand, category, portion, sunit, price, eff, ltype, ref_code, qty, unit) = v
        code_u = _norm(code).upper()
        row_errors = []
        if not code_u:
            _err(row_errors, row, 'Recipe Code', 'Mandatory field is blank', code)
        if not _norm(name):
            _err(row_errors, row, 'Recipe Name', 'Mandatory field is blank', name)
        brand_doc = m['brands'].get(_norm(brand).upper())
        if not _norm(brand):
            _err(row_errors, row, 'Brand Code', 'Mandatory field is blank', brand)
        elif not brand_doc:
            _err(row_errors, row, 'Brand Code', 'Brand code not found in Brand master', brand)
        lt = _norm(ltype).upper().replace('-', '').replace('_', '').replace(' ', '')
        if lt in ('SUBRECIPE', 'SUB'):
            lt = 'SUBRECIPE'
        elif lt in ('ITEM', 'RAW', 'RAWITEM'):
            lt = 'ITEM'
        else:
            _err(row_errors, row, 'Line Type', 'Must be ITEM or SUBRECIPE', ltype)
            lt = ''
        ref = None
        if not _norm(ref_code):
            _err(row_errors, row, 'Item / Sub Recipe Code', 'Mandatory field is blank', ref_code)
        elif lt == 'ITEM':
            ref = m['items'].get(_norm(ref_code).upper())
            if not ref:
                _err(row_errors, row, 'Item / Sub Recipe Code', 'Item code not found in Item master', ref_code)
        elif lt == 'SUBRECIPE':
            ref = m['subs'].get(_norm(ref_code).upper())
            if not ref:
                _err(row_errors, row, 'Item / Sub Recipe Code', 'Sub-recipe code not found. Upload sub-recipes first.', ref_code)
            elif not ref.get('active', True):
                _err(row_errors, row, 'Item / Sub Recipe Code', 'Sub-recipe is inactive', ref_code)
        q = _num(qty)
        if q is None or q <= 0:
            _err(row_errors, row, 'Qty Per Portion', 'Must be a number greater than zero', qty)
        if _norm(unit).upper() not in m['units']:
            _err(row_errors, row, 'Unit', 'Unit code not found in Unit master', unit)
        p = _num(price)
        if _norm(price) and (p is None or p < 0):
            _err(row_errors, row, 'Selling Price', 'Must be a number of zero or more', price)

        key = code_u or f"__row{row}"
        if key not in groups:
            if code_u and code_u in m['recipes']:
                _err(row_errors, row, 'Recipe Code', 'Recipe code already exists in the system', code)
            groups[key] = {'key': key, 'code': code_u, 'name': _norm(name),
                           'brand_id': (brand_doc or {}).get('id', ''),
                           'brand_name': (brand_doc or {}).get('name', ''),
                           'category': _norm(category), 'portion_size': _norm(portion),
                           'selling_unit': _norm(sunit) or 'Portion', 'selling_price': p or 0,
                           'effective_date': _date(eff), 'rows': [], 'items': [], 'errors': []}
            order.append(key)
        g = groups[key]
        if g['rows'] and g['name'] != _norm(name):
            _err(row_errors, row, 'Recipe Name', 'Header values differ from the first row of this recipe code', name)
        if ref and any(x['ref_id'] == ref['id'] for x in g['items']):
            _err(row_errors, row, 'Item / Sub Recipe Code', 'Same ingredient repeated within this recipe', ref_code)
        g['rows'].append(row)
        if row_errors:
            g['errors'] += row_errors
            errors += row_errors
        elif ref:
            g['items'].append({'line_type': lt, 'ref_id': ref['id'],
                               'name': ref.get('name', ''), 'qty': q, 'unit': _norm(unit).upper(), 'row': row})
    for key in order:
        g = groups[key]
        if not g['items'] and not g['errors']:
            _err(g['errors'], g['rows'][0], 'Item / Sub Recipe Code', 'Recipe has no valid ingredient rows', '')
            errors += g['errors']
        g['valid'] = not g['errors']
    return [groups[k] for k in order], errors


async def validate_sales(rows, m):
    groups, order, errors = {}, [], []
    for row, v in rows:
        date, brand, kitchen, recipe_code, sold, cancelled, comp = v
        row_errors = []
        d = _date(date)
        if not d:
            _err(row_errors, row, 'Date', 'Mandatory field is blank', date)
        elif len(d) != 10 or d[4] != '-' or d[7] != '-':
            _err(row_errors, row, 'Date', 'Use YYYY-MM-DD format', date)
        brand_doc = m['brands'].get(_norm(brand).upper())
        if not _norm(brand):
            _err(row_errors, row, 'Brand Code', 'Mandatory field is blank', brand)
        elif not brand_doc:
            _err(row_errors, row, 'Brand Code', 'Brand code not found in Brand master', brand)
        kdoc = m['kitchens'].get(_norm(kitchen).upper())
        if not _norm(kitchen):
            _err(row_errors, row, 'Kitchen Code', 'Mandatory field is blank', kitchen)
        elif not kdoc:
            _err(row_errors, row, 'Kitchen Code', 'Location code not found in Location master', kitchen)
        elif not kdoc.get('active', True):
            _err(row_errors, row, 'Kitchen Code', 'Location is inactive', kitchen)
        rec = m['recipes'].get(_norm(recipe_code).upper())
        if not _norm(recipe_code):
            _err(row_errors, row, 'Recipe Code', 'Mandatory field is blank', recipe_code)
        elif not rec:
            _err(row_errors, row, 'Recipe Code', 'Recipe code not found in Recipe master', recipe_code)
        elif not rec.get('active', True):
            _err(row_errors, row, 'Recipe Code', 'Recipe is inactive', recipe_code)
        elif brand_doc and rec.get('brand_id') != brand_doc['id']:
            _err(row_errors, row, 'Recipe Code', 'Recipe does not belong to the given brand', recipe_code)
        s = _num(sold)
        if s is None or s < 0:
            _err(row_errors, row, 'Qty Sold', 'Must be a number of zero or more', sold)
        c = _num(cancelled) or 0
        cp = _num(comp) or 0
        if c < 0 or cp < 0:
            _err(row_errors, row, 'Cancelled / Complimentary', 'Cannot be negative', cancelled)
        elif s is not None and (s - c < 0 or s - c - cp < 0):
            _err(row_errors, row, 'Cancelled / Complimentary', 'Cannot exceed quantity sold', cancelled)

        key = f"{d}|{(brand_doc or {}).get('id', _norm(brand))}|{(kdoc or {}).get('id', _norm(kitchen))}"
        if key not in groups:
            groups[key] = {'key': key, 'date': d, 'brand_id': (brand_doc or {}).get('id', ''),
                           'brand_name': (brand_doc or {}).get('name', ''),
                           'kitchen_id': (kdoc or {}).get('id', ''),
                           'kitchen_name': (kdoc or {}).get('name', ''),
                           'rows': [], 'items': [], 'errors': [], 'duplicate': False}
            order.append(key)
            if d and brand_doc and kdoc:
                existing = await db.daily_sales.find_one({'date': d, 'brand_id': brand_doc['id'],
                                                          'kitchen_id': kdoc['id'],
                                                          'archived': {'$ne': True},
                                                          'status': {'$in': ['draft', 'posted']}}, {'_id': 0})
                if existing:
                    groups[key]['duplicate'] = True
                    _err(row_errors, row, 'Date / Brand / Kitchen',
                         f"Sales already exist for this date, brand and kitchen ({existing['sale_no']}, {existing['status']}). Row blocked to prevent duplicate entry.", d)
        g = groups[key]
        if rec and any(x['recipe_id'] == rec['id'] for x in g['items']):
            _err(row_errors, row, 'Recipe Code', 'Same recipe repeated for this date/brand/kitchen', recipe_code)
        g['rows'].append(row)
        if row_errors:
            g['errors'] += row_errors
            errors += row_errors
        elif rec:
            g['items'].append({'recipe_id': rec['id'], 'recipe_name': rec['name'], 'recipe_code': rec['code'],
                               'qty_sold': s, 'cancelled': c, 'complimentary': cp,
                               'amount': round((s - c - cp) * (rec.get('selling_price') or 0), 2), 'row': row})
    for key in order:
        g = groups[key]
        if not g['items'] and not g['errors']:
            _err(g['errors'], g['rows'][0], 'Qty Sold', 'No valid sales rows for this date/brand/kitchen', '')
            errors += g['errors']
        g['total_amount'] = round(sum(i['amount'] for i in g['items']), 2)
        g['valid'] = not g['errors']
    return [groups[k] for k in order], errors


VALIDATORS = {'item': validate_item, 'subrecipe': validate_subrecipe, 'recipe': validate_recipe, 'sales': validate_sales}


@router.post('/{kind}/validate')
async def validate_upload(kind: str, file: UploadFile = File(...),
                          user=Depends(require_roles('admin', 'chef', 'store_manager'))):
    if kind not in SPECS:
        raise HTTPException(404, 'Unknown upload type')
    rows = await _read_rows(kind, file)
    groups, errors = await VALIDATORS[kind](rows, await _masters())
    valid = [g for g in groups if g['valid']]
    return {
        'kind': kind,
        'total_rows': len(rows),
        'total_groups': len(groups),
        'valid_groups': len(valid),
        'failed_groups': len(groups) - len(valid),
        'valid_rows': sum((len(g['rows']) if kind == 'item' else len(g['items'])) for g in valid),
        'failed_rows': len({e['row'] for e in errors}),
        'groups': groups,
        'errors': sorted(errors, key=lambda e: e['row']),
    }


class ImportIn(BaseModel):
    groups: list[dict]


@router.post('/{kind}/import')
async def import_upload(kind: str, body: ImportIn,
                        user=Depends(require_roles('admin', 'chef', 'store_manager'))):
    if kind not in SPECS:
        raise HTTPException(404, 'Unknown upload type')
    incoming = [g for g in body.groups if g.get('valid')]
    if not incoming:
        raise HTTPException(400, 'Nothing valid to import')
    imported, failed = [], []
    for g in incoming:
        rows = g.get('rows') or []
        first_row = rows[0] if rows else 0
        try:
            if kind == 'item':
                payload = ItemIn(code=g['code'], name=g['name'], category_id=g.get('category_id', ''),
                                 subcategory=g.get('subcategory', ''), purchase_unit=g['purchase_unit'],
                                 consumption_unit=g['consumption_unit'],
                                 conversion_factor=float(g.get('conversion_factor') or 1),
                                 min_stock=float(g.get('min_stock') or 0), reorder_level=float(g.get('reorder_level') or 0),
                                 max_stock=float(g.get('max_stock') or 0), rate=float(g.get('rate') or 0),
                                 supplier_id=g.get('supplier_id', ''))
                doc = await create_item(payload, user)
                detail = f"{doc['purchase_unit']} / {doc['consumption_unit']}"
                op = g.get('opening')
                if op:
                    qb, _ = await to_base(float(op['qty']), op['unit'])
                    await post_txn(doc['id'], qb, 'opening', to_location_id=op['location_id'], user=user,
                                   qty=float(op['qty']), unit=op['unit'], reference_no='OPENING',
                                   rate_per_base=doc.get('rate_per_base') or 0, remarks='Bulk opening stock')
                    detail += f" · opening {op['qty']} {op['unit']} @ {op['location_name']}"
                imported.append({'ref': doc['code'], 'label': doc['name'], 'rows': rows, 'detail': detail})
            elif kind == 'subrecipe':
                payload = SubRecipeIn(name=g['name'], code=g['code'], brand_id=g.get('brand_id') or '',
                                      output_item_id=g['output_item_id'], yield_qty=float(g['yield_qty']),
                                      yield_unit=g['yield_unit'],
                                      items=[SubRecipeItemIn(item_id=i['item_id'], qty=float(i['qty']), unit=i['unit'])
                                             for i in g['items']])
                doc = await create_subrecipe(payload, user)
                imported.append({'ref': doc['code'], 'label': doc['name'], 'rows': rows,
                                 'detail': f"{len(g['items'])} ingredients, yield {doc['yield_pct']}%"})
            elif kind == 'recipe':
                payload = RecipeIn(name=g['name'], code=g['code'], brand_id=g['brand_id'],
                                   category=g.get('category', ''), portion_size=g.get('portion_size', ''),
                                   selling_unit=g.get('selling_unit') or 'Portion',
                                   selling_price=float(g.get('selling_price') or 0),
                                   effective_date=g.get('effective_date') or '',
                                   items=[RecipeItemIn(
                                       item_id=i['ref_id'] if i['line_type'] == 'ITEM' else '',
                                       sub_recipe_id=i['ref_id'] if i['line_type'] == 'SUBRECIPE' else '',
                                       qty=float(i['qty']), unit=i['unit']) for i in g['items']])
                doc = await create_recipe(payload, user)
                imported.append({'ref': doc['code'], 'label': doc['name'], 'rows': rows,
                                 'detail': f"{len(g['items'])} ingredient lines, v1"})
            else:
                payload = SaleIn(date=g['date'], brand_id=g['brand_id'], kitchen_id=g['kitchen_id'],
                                 source='bulk_upload',
                                 items=[SaleItemIn(recipe_id=i['recipe_id'], qty_sold=float(i['qty_sold']),
                                                   cancelled=float(i['cancelled']),
                                                   complimentary=float(i['complimentary'])) for i in g['items']])
                doc = await create_sale(payload, user)
                imported.append({'ref': doc['sale_no'], 'label': f"{g['date']} · {g.get('kitchen_name', '')}",
                                 'rows': rows, 'detail': f"{len(g['items'])} products, draft total {doc['total_amount']}"})
        except HTTPException as e:
            failed.append({'ref': g.get('code') or g.get('date', ''), 'rows': rows,
                           'row': first_row, 'reason': str(e.detail)})
        except Exception as e:
            failed.append({'ref': g.get('code') or g.get('date', ''), 'rows': rows,
                           'row': first_row, 'reason': f'Unexpected error: {e}'})
    return {
        'kind': kind,
        'total_groups': len(incoming),
        'imported_groups': len(imported),
        'failed_groups': len(failed),
        'imported_rows': sum(len(i['rows']) for i in imported),
        'failed_rows': sum(len(f['rows']) for f in failed),
        'imported': imported,
        'failed': failed,
    }
