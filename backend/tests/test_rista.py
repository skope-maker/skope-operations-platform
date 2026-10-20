"""Rista file imports and reports, using small made-up files shaped like the real Rista downloads.

Run against a backend on a throwaway database with no Rista data loaded yet (the totals are exact):
    RISTA_API=http://localhost:8011 pytest backend/tests/test_rista.py -n 0
"""
import csv
import io
import os

import pytest
import requests
from openpyxl import Workbook, load_workbook

API = os.environ.get('RISTA_API', 'http://localhost:8011').rstrip('/') + '/api'
ADMIN = (os.environ.get('E2E_ADMIN_EMAIL', 'chennakeshavaediga@gmail.com'), os.environ.get('E2E_ADMIN_PASSWORD', 'admin123'))
DEMO_PASSWORD = os.environ.get('E2E_DEMO_PASSWORD', 'demo123')

SALES_HEADERS = ['Branch Name', 'Branch Code', 'Business Brand', 'Business Date', 'Invoice Date', 'Invoice Number',
                 'Order Source', 'Source Order Number', 'Invoice Type', 'Sale Status', 'SKU', 'Type', 'Item Name',
                 'Sale Item Qty', 'Unit Rate', 'Gross Amount', 'Discounts', 'Net Amount', 'Taxes',
                 'Tax Collected by Aggregator', 'Direct Charge', 'Indirect Charge', 'Channel', 'Account Name',
                 'Customer ID', 'Customer Phone', 'Customer Name', 'Punched Name', 'Notes']


def sale(branch, code, brand, day, inv, source, sku, typ, name, qty, gross, disc, channel, account):
    return [branch, code, brand, day, f'{day} 12:00:00', inv, source, f'SRC{inv}', 'Sale', 'Closed', sku, typ, name,
            qty, gross / qty if qty else 0, gross, -disc, gross - disc, 0, 0, 0, 0, channel, account,
            'CUST1', '9999999999', 'Secret Customer', 'Staff Person', 'call me']


SALES = [
    sale('JP Nagar', 'BEN', 'Kochi Kurry Klub', '2026-09-20', '100', 'zomato', 'KKK1', 'Item', 'Appam', 2, 200, 20, 'Zomato(Kochi Kurry Klub)', 'Kochi Kurry Klub'),
    sale('JP Nagar', 'BEN', 'Kochi Kurry Klub', '2026-09-20', '100', 'zomato', 'ADD1', 'Add On', 'Extra Curry', 1, 50, 0, 'Zomato(Kochi Kurry Klub)', 'Kochi Kurry Klub'),
    sale('JP Nagar', 'BEN', 'Kochi Kurry Klub', '2026-09-20', '101', 'swiggy', 'SHW1', 'Item', 'Shawarma', 1, 150, 0, 'Swiggy(Kochi Kurry Klub)', 'Al Mashawi Shawarma'),
    sale('Marathahalli', 'MAR', 'Al Mashawi Shawarma', '2026-09-21', '200', 'swiggy', 'SHW1', 'Item', 'Shawarma', 3, 450, 50, 'Swiggy (Al-Mashawi)', 'Al Mashawi Shawarma'),
    sale('Marathahalli', 'MAR', 'Al Mashawi Shawarma', '2026-09-21', '201', 'POS', 'NEW9', 'Item', 'Not On Menu', 1, 100, 0, 'Offline - Marathahalli', 'Al Mashawi Shawarma'),
]
# totals: 4 orders on 2 days; net = 180 + 50 + 150 + 400 + 100 = 880

CANCELLED_HEADERS = ['Branch Name', 'Branch Code', 'Business Brand', 'Business Date', 'Invoice Number', 'Order Number',
                     'Order Source', 'Source Order Number', 'Order/Invoice Date', 'Status', 'Fulfillment Status',
                     'Channel', 'Time to Cancel', 'Cancel/Void Date', 'Cancel/Void By', 'Reason For Cancel/Void',
                     'Remarks', 'Source Reason for Cancel/Void', 'Gross Amount', 'Discounts', 'Net Amount', 'Taxes',
                     'Total', 'Customer Name', 'Customer Phone']
CANCELLED = [
    ['JP Nagar', 'BEN', 'Kochi Kurry Klub', '2026-09-20', '900', '', 'zomato', 'Z1', '2026-09-20 20:00:00', 'Voided',
     'Accepted', 'Zomato(Kochi Kurry Klub)', '5.0', '2026-09-20 20:05:00', 'SYSTEM', 'Cancelled by zomato. reject',
     'remark', "Order rejected due to 'Customer Cancellation'", 300, -30, 270, 0, 270, 'Secret Customer', '9999999999'],
    ['Marathahalli', 'MAR', 'Al Mashawi Shawarma', '2026-09-21', '901', '', 'POS', '', '2026-09-21 13:00:00', 'Cancelled',
     'Created', 'Offline - Marathahalli', '1.0', '2026-09-21 13:01:00', 'Staff (X1)', 'Wrongly Punched', '', '',
     100, 0, 100, 0, 100, '', ''],
]

MENU_HEADERS = ['SKU', 'MASTER CODE', 'CATEGORY', 'ITEM NAME', 'VARIANT ATTRIBUTES', 'ITEM BRAND', 'STATUS', 'SUB CATEGORY',
                'ACCOUNT', 'MEASURING UNIT', 'PRICE INCLUDES TAX', 'MRP', 'Price-JP Nagar:Swiggy(Kochi Kurry Klub)',
                'Price-JP Nagar:Zomato(Kochi Kurry Klub)', 'Price-JP Nagar:Offline - JP Nagar', 'Price-Marathahalli Group:Swiggy (Al-Mashawi)',
                'Branch Code Deactivation :All Channels', 'TAGS', 'DESCRIPTION']
MENU = [
    ['KKK1', '', 'Breakfast', 'Appam', '', '', 'Active', '', 'Kochi Kurry Klub', 'ea', 'No', '', '120', '120', '99', '', '', 'Vegetarian', 'Soft appam'],
    ['shw1', '', 'Rolls', 'Shawarma', '', '', 'Active', '', 'Al Mashawi Shawarma', 'ea', 'No', '', '', '', '', '150', '', 'Non-vegetarian', ''],
    ['OLD1', '', 'Rolls', 'Old Roll', '', '', 'Inactive', '', 'Al Mashawi Shawarma', 'ea', 'No', '', '', '', '', '', '', '', ''],
]


def csv_bytes(headers, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return ('﻿' + buf.getvalue()).encode('utf-8')


def xlsx_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def session(email, password):
    r = requests.post(f'{API}/auth/login', json={'email': email, 'password': password}, timeout=30)
    assert r.status_code == 200, f'login {email}: {r.status_code} {r.text[:200]}'
    s = requests.Session()
    s.headers['Authorization'] = f"Bearer {r.json()['token']}"
    return s


@pytest.fixture(scope='module')
def admin():
    return session(*ADMIN)


@pytest.fixture(scope='module')
def management():
    return session('management@kitchen.com', DEMO_PASSWORD)


@pytest.fixture(scope='module')
def chef():
    return session('chef@kitchen.com', DEMO_PASSWORD)


def upload(s, kind, content, name):
    return s.post(f'{API}/rista/import/{kind}', files={'file': (name, content)}, timeout=60)


@pytest.fixture(scope='module')
def loaded(admin):
    r = upload(admin, 'sales', csv_bytes(SALES_HEADERS, SALES), 'Multidate - Sales Items By Transaction.csv')
    assert r.status_code == 200, r.text
    r2 = upload(admin, 'cancelled', csv_bytes(CANCELLED_HEADERS, CANCELLED), 'Multidate - Cancelled Transactions.csv')
    assert r2.status_code == 200, r2.text
    r3 = upload(admin, 'menu', csv_bytes(MENU_HEADERS, MENU), 'Catalog.csv')
    assert r3.status_code == 200, r3.text
    return r.json(), r2.json(), r3.json()


def test_sales_import_summary(loaded):
    s, _, _ = loaded
    assert s['rows'] == 5 and s['orders'] == 4 and s['net'] == 880
    assert s['date_from'] == '2026-09-20' and s['date_to'] == '2026-09-21' and s['days'] == 2
    assert s['branches'] == ['JP Nagar', 'Marathahalli']


def test_customer_details_are_not_stored(admin, loaded):
    r = admin.get(f'{API}/rista/sales/export?format=csv', timeout=30)
    assert r.status_code == 200
    body = r.content.decode('utf-8-sig')
    for secret in ('Secret Customer', '9999999999', 'Staff Person', 'call me', 'CUST1'):
        assert secret not in body
    c = admin.get(f'{API}/rista/cancelled', timeout=30).json()
    assert all('customer_name' not in row and 'customer_phone' not in row for row in c['rows'])
    assert 'Secret Customer' not in str(c)


def test_reimport_replaces_instead_of_doubling(admin, loaded):
    again = upload(admin, 'sales', csv_bytes(SALES_HEADERS, SALES), 'again.csv').json()
    assert again['replaced'] == 5
    t = admin.get(f'{API}/rista/sales/summary', timeout=30).json()['totals']
    assert t['orders'] == 4 and t['lines'] == 5 and t['net'] == 880


def test_summary_totals_and_breakdowns(admin, loaded):
    d = admin.get(f'{API}/rista/sales/summary', timeout=30).json()
    t = d['totals']
    assert t['gross'] == 950 and t['discounts'] == -70 and t['net'] == 880
    assert t['qty'] == 7  # items only, the add-on is not counted
    assert t['days'] == 2 and t['avg_order'] == 220
    days = {r['key']: r for r in d['by_day']}
    assert days['2026-09-20']['orders'] == 2 and days['2026-09-20']['net'] == 380
    brands = {r['key']: r for r in d['by_brand']}
    assert brands['Kochi Kurry Klub']['net'] == 380 and brands['Al Mashawi Shawarma']['net'] == 500
    sources = {r['key']: r for r in d['by_source']}
    assert sources['swiggy']['orders'] == 2 and sources['swiggy']['net'] == 550
    top = {r['sku']: r for r in d['top_dishes']}
    assert top['SHW1']['qty'] == 4 and 'ADD1' not in top


def test_filters_and_brand_field(admin, loaded):
    by_business = admin.get(f'{API}/rista/sales/summary', params={'brand': 'Al Mashawi Shawarma'}, timeout=30).json()['totals']
    assert by_business['net'] == 500 and by_business['orders'] == 2
    by_account = admin.get(f'{API}/rista/sales/summary', params={'brand': 'Al Mashawi Shawarma', 'brand_field': 'account_name'}, timeout=30).json()['totals']
    assert by_account['net'] == 650 and by_account['orders'] == 3  # includes the KKK order's shawarma line
    one_day = admin.get(f'{API}/rista/sales/summary', params={'date_from': '2026-09-21', 'date_to': '2026-09-21', 'branch': 'MAR'}, timeout=30).json()['totals']
    assert one_day['net'] == 500
    bad = admin.get(f'{API}/rista/sales/summary', params={'brand_field': 'customer_name'}, timeout=30)
    assert bad.status_code == 400
    f = admin.get(f'{API}/rista/sales/filters', timeout=30).json()
    assert {'code': 'BEN', 'name': 'JP Nagar'} in f['branches'] and f['date_from'] == '2026-09-20'
    assert 'Al Mashawi Shawarma' in f['account_names'] and 'pos' not in f['sources']


def test_export_xlsx(admin, loaded):
    r = admin.get(f'{API}/rista/sales/export', params={'format': 'xlsx', 'branch': 'BEN'}, timeout=30)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == 'Business Date' and len(rows) == 4
    assert admin.get(f'{API}/rista/sales/export', params={'format': 'pdf'}, timeout=30).status_code == 400


def test_cancelled(admin, loaded):
    _, c, _ = loaded
    assert c['rows'] == 2 and c['total'] == 370
    d = admin.get(f'{API}/rista/cancelled', timeout=30).json()
    assert d['totals'] == {'orders': 2, 'total': 370, 'by_platform': 1, 'by_staff': 1}
    assert d['by_reason'][0]['orders'] == 1
    only_mar = admin.get(f'{API}/rista/cancelled', params={'branch': 'MAR'}, timeout=30).json()
    assert only_mar['totals']['orders'] == 1 and only_mar['rows'][0]['reason'] == 'Wrongly Punched'
    x = admin.get(f'{API}/rista/cancelled/export', params={'format': 'csv'}, timeout=30)
    assert x.status_code == 200 and 'Wrongly Punched' in x.text and 'Secret Customer' not in x.text


def test_menu_and_checks(admin, loaded):
    _, _, m = loaded
    assert m['rows'] == 3 and m['active'] == 2 and m['accounts'] == 2 and m['price_columns'] == 4
    rows = {r['sku']: r for r in admin.get(f'{API}/rista/menu', timeout=30).json()}
    assert rows['KKK1']['price_min'] == 99 and rows['KKK1']['price_max'] == 120 and rows['KKK1']['sold_qty'] == 2
    assert rows['SHW1']['sold_qty'] == 4 and rows['OLD1']['active'] is False
    active = admin.get(f'{API}/rista/menu', params={'status': 'active', 'q': 'appam'}, timeout=30).json()
    assert [r['sku'] for r in active] == ['KKK1']
    one = admin.get(f'{API}/rista/menu/kkk1', timeout=30).json()
    assert {'branch': 'JP Nagar', 'channel': 'Offline - JP Nagar', 'price': 99} in one['prices']
    checks = admin.get(f'{API}/rista/menu/checks', timeout=30).json()
    assert checks['menu_loaded'] and checks['sold_dishes'] == 3
    assert [r['sku'] for r in checks['not_in_menu']] == ['NEW9']
    assert {r['sku'] for r in checks['without_recipe']} == {'KKK1', 'SHW1', 'NEW9'}
    assert admin.get(f'{API}/rista/menu/accounts', timeout=30).json() == ['Al Mashawi Shawarma', 'Kochi Kurry Klub']


def test_excel_download_is_accepted(admin, loaded):
    r = upload(admin, 'cancelled', xlsx_bytes(CANCELLED_HEADERS, CANCELLED), 'cancelled.xlsx')
    assert r.status_code == 200, r.text
    assert r.json()['rows'] == 2 and r.json()['replaced'] == 2


def test_wrong_or_empty_file_is_refused(admin):
    r = upload(admin, 'sales', csv_bytes(MENU_HEADERS, MENU), 'Catalog.csv')
    assert r.status_code == 400 and 'not the Rista' in r.json()['detail']
    assert upload(admin, 'menu', b'', 'empty.csv').status_code == 400
    assert upload(admin, 'sales', csv_bytes(SALES_HEADERS, []), 'header-only.csv').status_code == 400


def test_import_history(admin, loaded):
    rows = admin.get(f'{API}/rista/imports', timeout=30).json()
    kinds = {r['kind'] for r in rows}
    assert {'sales', 'cancelled', 'menu'} <= kinds
    assert rows[0]['imported_by']


def test_permissions(management, chef, loaded):
    assert management.get(f'{API}/rista/sales/summary', timeout=30).status_code == 200
    assert management.get(f'{API}/rista/menu', timeout=30).status_code == 200
    assert upload(management, 'sales', csv_bytes(SALES_HEADERS, SALES), 'x.csv').status_code == 403
    for path in ('/rista/sales/summary', '/rista/cancelled', '/rista/menu', '/rista/imports', '/rista/sales/export'):
        assert chef.get(f'{API}{path}', timeout=30).status_code == 403, path
    assert requests.get(f'{API}/rista/sales/summary', timeout=30).status_code == 401
