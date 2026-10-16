import io, json, os, requests
from openpyxl import load_workbook

API = (os.environ.get("TEST_BASE") or os.environ.get("REACT_APP_BACKEND_URL") or "http://localhost:8001").rstrip("/")
tok = requests.post(f"{API}/api/auth/login", json={"email": "chennakeshavaediga@gmail.com", "password": "admin123"}).json()["token"]
H = {"Authorization": f"Bearer {tok}"}

# pick a real location code
kitchens = requests.get(f"{API}/api/kitchens", headers=H).json()
store = next((k for k in kitchens if k["type"] == "store"), kitchens[0])
loc_code = store["code"]
print("Using location code:", loc_code, store["name"])

# build item file from the template layout
r = requests.get(f"{API}/api/bulk/template/item", headers=H)
wb = load_workbook(io.BytesIO(r.content))
ws = wb["Items"]
# clear example rows (rows 2..end)
ws.delete_rows(2, ws.max_row)
# valid item with opening stock
ws.append(["ITM-BULKTEST1", "Bulk Test Onion", "", "", "KG", "GM", 1000, 5000, 10000, 50000, 40, "", 25, "KG", loc_code])
# valid item no opening
ws.append(["ITM-BULKTEST2", "Bulk Test Salt", "", "", "KG", "GM", 1000, 2000, 5000, 20000, 12, "", "", "", ""])
# invalid: bad unit
ws.append(["ITM-BULKTEST3", "Bad Unit Item", "", "", "XX", "GM", 1000, 0, 0, 0, 0, "", "", "", ""])
# invalid: duplicate code
ws.append(["ITM-BULKTEST1", "Dup Code", "", "", "KG", "GM", 1000, 0, 0, 0, 0, "", "", "", ""])
buf = io.BytesIO(); wb.save(buf); buf.seek(0)

v = requests.post(f"{API}/api/bulk/item/validate", headers=H, files={"file": ("items.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()
print("VALIDATE: total=%s valid=%s failed=%s valid_rows=%s" % (v["total_groups"], v["valid_groups"], v["failed_groups"], v["valid_rows"]))
assert v["valid_groups"] == 2, v
assert v["failed_groups"] == 2, [e for e in v["errors"]]

imp = requests.post(f"{API}/api/bulk/item/import", headers=H, json={"groups": [g for g in v["groups"] if g["valid"]]}).json()
print("IMPORT: imported=%s failed=%s" % (imp["imported_groups"], imp["failed_groups"]))
assert imp["imported_groups"] == 2, imp
for i in imp["imported"]:
    print("  ->", i["ref"], "|", i["detail"])

# verify opening stock landed at store
stock = requests.get(f"{API}/api/inventory/stock?location_id={store['id']}", headers=H).json()
onion = next((s for s in stock if s["item_code"] == "ITM-BULKTEST1"), None)
print("Opening stock ITM-BULKTEST1:", onion["stock"], onion["unit"] if onion else None)
assert onion and onion["stock"] == 25000, onion  # 25 KG = 25000 GM

# reorder list should now show BULKTEST items (0 stock, min>0)
reo = requests.get(f"{API}/api/reports/reorder", headers=H).json()
codes = [r["item_code"] for r in reo]
print("Reorder list codes:", codes[:10])
assert "ITM-BULKTEST2" in codes, "salt with 0 stock and min should appear"

# cleanup: delete the test items
items = requests.get(f"{API}/api/items", headers=H).json()
for it in items:
    if it["code"] in ("ITM-BULKTEST1", "ITM-BULKTEST2"):
        requests.delete(f"{API}/api/items/{it['id']}", headers=H)
print("ALL ITEM-BULK CHECKS PASSED")
